"""Phase 4 contracts: real local parsers/OCR plus isolated provider/network cases."""
import asyncio
import json
import os
import sqlite3
import time
import zipfile
from pathlib import Path
import httpx
import pytest
from docx import Document
from PIL import Image,ImageDraw,ImageFont
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
from fastapi import FastAPI
from fastapi.testclient import TestClient
from server.personal_store import PersonalStore,PersonalError,encode
from server.review_store import ReviewStore
from server.extraction_models import ConfirmationInput
from server.extraction_provider import extraction_messages
from server.extraction_validation import validate_candidate,parse_output
from server.extraction_service import ExtractionService
from server.extraction_errors import ExtractionFailure
from server.documents.policy import DocumentError,validate_file,MAX_BYTES
from server.documents.parsers import parse_document,native_ocr
from server.documents.storage import AttachmentStore
from server.documents.service import DocumentService
from server.documents.blocks import span_source,select_model_blocks,document_warnings
from server.documents.downloader import download
from server.documents.api import router,UploadBodyLimit
from server.review_api import router as review_router
from server.test_extraction_review import candidate,FixtureProvider,TestSettings,finish


@pytest.fixture
def store(tmp_path):
    s=PersonalStore(tmp_path/'personal.sqlite3');s.initialize();return s


def text_pdf(path,texts):
    writer=PdfWriter()
    for text in texts:
        page=writer.add_blank_page(width=612,height=792)
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
        content=DecodedStreamObject();content.set_data(('BT /F1 16 Tf 40 700 Td ('+text+') Tj ET').encode('ascii'))
        page[NameObject('/Contents')]=writer._add_object(content)
    writer.write(path);return path


def local_snapshot(store,path,body='',mime=None,url=None):
    files=AttachmentStore(store)
    raw=files.collect([(path.name,mime,path)],body,'文件测试')
    att=raw['snapshot']['attachments'][0]
    parsed=parse_document(path,path.suffix[1:],path.parent)
    meta=parsed['metadata']
    with store.connection(write=True) as c:
        if url:c.execute('UPDATE attachment SET original_url=? WHERE id=?',(url,att['id']))
        c.execute("UPDATE attachment SET parse_status='parsed',extracted_text=?,extractor=?,extractor_version=?,parse_manifest_json=?,needs_review=? WHERE id=?",(parsed['extracted_text'],meta['extractor'],meta['extractor_version'],encode(parsed),int(meta['needs_review']),att['id']))
    return ReviewStore(store).get(files.compose(raw['snapshot']['id'],[att['id']]))


def review_payload(raw,block,quote,value='2026-10-18',ack=True):
    span={'block_id':block['block_id'],'quote':quote,'start_offset':block['text'].index(quote),'end_offset':block['text'].index(quote)+len(quote)}
    return ConfirmationInput(expected_version=raw['version'],acknowledged_warnings=ack,events=[{'id':'e1','title':'附件报名验收','participation':'yes','relevance':'yes',
      'deadlines':[{'id':'d1','role':'报名截止','value':value,'precision':'date','original_text':quote,'evidence':[span],'confirmed_by_user':True}],
      'actions':[{'id':'a1','title':'提交材料','kind':'explicit','deadline_id':'d1','evidence':[span]}]}])


def test_case01_text_pdf_deadline_and_page_evidence(store,tmp_path):
    raw=local_snapshot(store,text_pdf(tmp_path/'notice.pdf',['Submit application before 2026-10-18.']))
    b=raw['snapshot']['blocks'][0]
    event=ReviewStore(store).confirm(raw['snapshot']['id'],review_payload(raw,b,'2026-10-18'))['events'][0]
    assert len(event['deadlines'])==len(event['tasks'])==1
    assert all(e['page']==1 and e['source_type']=='pdf' and e['attachment_id'] for e in event['evidence'])
    assert Path(AttachmentStore(store).get(b['attachment_id'])['local_path']).is_file()


def test_case02_multipage_independent_evidence(store,tmp_path):
    raw=local_snapshot(store,text_pdf(tmp_path/'multiple.pdf',['Submit application before 2026-10-18.','Submit application before 2026-10-19.']))
    blocks=raw['snapshot']['blocks'];assert [b['page'] for b in blocks]==[1,2]
    p=review_payload(raw,blocks[0],'2026-10-18').model_dump()
    d={**p['events'][0]['deadlines'][0],'id':'d2','value':'2026-10-19','original_text':'2026-10-19','evidence':[{'block_id':blocks[1]['block_id'],'quote':'2026-10-19','start_offset':33,'end_offset':43}]}
    d['evidence'][0]['start_offset']=blocks[1]['text'].index('2026-10-19');d['evidence'][0]['end_offset']=d['evidence'][0]['start_offset']+10
    p['events'][0]['deadlines'].append(d)
    event=ReviewStore(store).confirm(raw['snapshot']['id'],ConfirmationInput.model_validate(p))['events'][0]
    assert len(event['deadlines'])==2 and {e['page'] for e in event['evidence']}=={1,2}


def screenshot(path,text='请于2026年10月18日前提交材料。'):
    im=Image.new('RGB',(1300,300),'white');font=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',44)
    ImageDraw.Draw(im).text((40,60),text,font=font,fill='black');im.save(path);return path


@pytest.mark.skipif(os.name!='nt',reason='Windows native OCR')
def test_case03_scanned_pdf_real_ocr(tmp_path):
    png=screenshot(tmp_path/'scan.png')
    with Image.open(png) as im: im.save(tmp_path/'scan.pdf','PDF',resolution=120)
    parsed=parse_document(tmp_path/'scan.pdf','pdf',tmp_path)
    assert parsed['metadata']['pdf_kind']=='scanned_pdf' and parsed['metadata']['needs_ocr']
    assert any('10月18' in b['text'] and b['page']==1 for b in parsed['blocks'])
    assert parsed['metadata']['needs_review']


def test_case04_docx_paragraphs_heading_order(tmp_path):
    path=tmp_path/'paragraph.docx';doc=Document();doc.add_heading('学校申报通知',1);doc.add_paragraph('请于2026年10月18日前提交。');doc.save(path)
    parsed=parse_document(path,'docx',tmp_path)
    assert parsed['blocks'][0]['text']=='学校申报通知' and 'Heading 1' in parsed['blocks'][0]['section']
    assert '段落 2' in parsed['blocks'][1]['section']


def test_case05_docx_table_date_evidence(store,tmp_path):
    path=tmp_path/'table.docx';doc=Document();doc.add_paragraph('报名安排');table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='截止日期';table.cell(0,1).text='2026年10月18日前提交';doc.save(path)
    raw=local_snapshot(store,path)
    b=next(b for b in raw['snapshot']['blocks'] if '10月18' in b['text'])
    event=ReviewStore(store).confirm(raw['snapshot']['id'],review_payload(raw,b,'2026年10月18日前'))['events'][0]
    assert '表格 1 · 行 1 · 列 2' in event['evidence'][0]['section'] and event['evidence'][0]['source_type']=='docx'


@pytest.mark.skipif(os.name!='nt',reason='Windows native OCR')
def test_case06_screenshot_ocr_original_bounding_box(store,tmp_path):
    raw=local_snapshot(store,screenshot(tmp_path/'screen.png'))
    b=next(b for b in raw['snapshot']['blocks'] if '10月18' in b['text'])
    e=span_source(raw['snapshot'],{'block_id':b['block_id'],'quote':'10月18'})
    assert e['bounding_box']['width']<b['bounding_box']['width'] and e['needs_review']
    event=ReviewStore(store).confirm(raw['snapshot']['id'],review_payload(raw,b,'2026年10月18日前'))['events'][0]
    assert event['evidence'][0]['source_type']=='image_ocr' and event['evidence'][0]['bounding_box']


@pytest.mark.skipif(os.name!='nt',reason='Windows native OCR')
def test_case07_low_ocr_requires_acknowledgement(store,tmp_path):
    path=screenshot(tmp_path/'small.png','10月18日')
    parsed=parse_document(path,'png',tmp_path);assert parsed['metadata']['ocr']['quality']=='low'
    raw=local_snapshot(store,path);b=raw['snapshot']['blocks'][0]
    assert raw['snapshot']['document_warnings'][0]['code']=='ocr_unverified'
    with pytest.raises(PersonalError,match='核对'): ReviewStore(store).confirm(raw['snapshot']['id'],review_payload(raw,b,b['text'],ack=False))
    with store.connection() as c: assert c.execute('SELECT COUNT(*) FROM event').fetchone()[0]==0


def test_case08_body_pdf_joint_input_same_item(store,tmp_path):
    raw=local_snapshot(store,text_pdf(tmp_path/'joint.pdf',['Submit application before 2026-10-18.']),'具体安排见附件，请提交材料。')
    assert {b['source_type'] for b in raw['snapshot']['blocks']}=={'text','pdf'}
    payload=json.loads(extraction_messages(raw['snapshot'])[1]['content'].split('只分析以下 JSON 中的 original_text 或 source_blocks 数据：\n')[1])
    assert payload['original_text']=='' and len(payload['source_blocks'])==2
    with store.connection() as c: assert c.execute('SELECT COUNT(*) FROM raw_item').fetchone()[0]==1 and c.execute('SELECT COUNT(*) FROM event').fetchone()[0]==0


def test_case09_body_attachment_conflict(store,tmp_path):
    path=tmp_path/'dates.txt';path.write_text('请于10月18日前提交报名材料。',encoding='utf-8')
    raw=local_snapshot(store,path,'请于10月15日前提交报名材料。')
    warnings=raw['snapshot']['document_warnings'];assert any(w['code']=='date_conflict' and len(w['sources'])==2 for w in warnings)
    b=raw['snapshot']['blocks'][-1]
    with pytest.raises(PersonalError,match='核对'): ReviewStore(store).confirm(raw['snapshot']['id'],review_payload(raw,b,'10月18日',ack=False))


def test_case10_same_name_content_new_version_preserves_old(store,tmp_path):
    path=tmp_path/'notice.txt';path.write_text('2026年10月18日前提交',encoding='utf-8');raw=local_snapshot(store,path,url='https://www.zqu.edu.cn/notice.txt');old=raw['snapshot']['attachments'][0];b=raw['snapshot']['blocks'][0]
    event=ReviewStore(store).confirm(raw['snapshot']['id'],review_payload(raw,b,'2026年10月18日前'))['events'][0]
    path.write_text('2026年10月19日前提交',encoding='utf-8');files=AttachmentStore(store);new=files.downloaded(old['id'],path,'text/plain')
    assert new['id']!=old['id'] and new['sha256']!=old['sha256'] and new['version']==2
    assert files.path(files.get(old['id'])).read_text(encoding='utf-8')=='2026年10月18日前提交'
    assert store.get_event(event['id'])['evidence']==event['evidence']
    with pytest.raises(sqlite3.IntegrityError,match='immutable'):
        with store.connection(write=True) as c:c.execute("UPDATE attachment SET filename='changed' WHERE id=?",(old['id'],))


def test_case11_oversize_stream_and_declared(store,tmp_path):
    path=tmp_path/'big.txt'
    with path.open('wb') as f:f.truncate(MAX_BYTES+1)
    with pytest.raises(DocumentError,match='50 MB'): AttachmentStore(store).preserve(path,'txt')
    app=FastAPI();app.add_middleware(UploadBodyLimit)
    with TestClient(app) as client: assert client.post('/api/personal/collect-files',headers={'content-length':'300000000'},content=b'a').status_code==413
    def fake(req):return httpx.Response(200,headers={'content-length':'300000000','content-type':'application/pdf'},content=b'fake')
    with pytest.raises(DocumentError) as e:download('https://www.zqu.edu.cn/test.pdf',['www.zqu.edu.cn'],tmp_path/'download',transport=httpx.MockTransport(fake),resolver=lambda _:None)
    assert e.value.code=='file_too_large'


def test_case12_wrong_mime_and_executable_magic(tmp_path):
    path=text_pdf(tmp_path/'mime.pdf',['School notice deadline 2026-10-18.'])
    with pytest.raises(DocumentError) as e:validate_file(path,'pdf','text/html')
    assert e.value.code=='mime_mismatch'
    path.write_bytes(b'MZ executable data')
    with pytest.raises(DocumentError):validate_file(path,'pdf','application/pdf')


@pytest.mark.parametrize('ext,data',[('pdf',b'%PDF-1.7\n corrupt'),('docx',b'PK corrupt')])
def test_case13_corrupt_document_preserved(store,tmp_path,ext,data):
    path=tmp_path/('bad.'+ext);path.write_bytes(data);files=AttachmentStore(store);raw=files.collect([(path.name,None,path)],'','坏文件')
    with pytest.raises(DocumentError):parse_document(path,ext,tmp_path)
    assert files.path(files.get(raw['snapshot']['attachments'][0]['id'])).read_bytes()==data


def test_case14_fabricated_block_rejected(store,tmp_path):
    raw=local_snapshot(store,text_pdf(tmp_path/'evidence.pdf',['Submit application before 2026-10-18.']))
    b=raw['snapshot']['blocks'][0];p=review_payload(raw,b,'2026-10-18').model_dump();p['events'][0]['deadlines'][0]['evidence'][0]['block_id']='ATTACHMENT_NONEXISTENT_PAGE_3'
    with pytest.raises(PersonalError,match='来源块'):ReviewStore(store).confirm(raw['snapshot']['id'],ConfirmationInput.model_validate(p))
    output=candidate(b['text'],[('2026-10-18','2026-10-18')],['Submit application'])
    for e in output['evidence']:e['block_id']='ATTACHMENT_NONEXISTENT_PAGE_3'
    assert all(not e['valid'] for e in validate_candidate(output,raw['snapshot'])['evidence'])
    with store.connection() as c:assert c.execute('SELECT COUNT(*) FROM event').fetchone()[0]==0


def test_case15_provider_failure_keeps_full_local_text(store,tmp_path):
    raw=local_snapshot(store,text_pdf(tmp_path/'offline.pdf',['Submit application before 2026-10-18.']))
    async def run():
        service=ExtractionService(store,TestSettings(),FixtureProvider(failure=ExtractionFailure('connection_error','测试离线')))
        await service.start(raw['snapshot']['id']);await finish(service);await service.close()
    asyncio.run(run());saved=ReviewStore(store).get(raw['snapshot']['id'])
    assert saved['extraction']['status']=='failed' and saved['snapshot']['blocks'][0]['text']==raw['snapshot']['blocks'][0]['text']
    assert AttachmentStore(store).get(saved['snapshot']['attachments'][0]['id'])['extracted_text']


def test_200page_input_budget_and_immutable_full_text():
    blocks=[{'block_id':f'p{i}','ordinal':i,'text':'普通说明。'*499,'attachment_id':'a','source_type':'pdf','page':i+1,'needs_review':False} for i in range(200)]
    blocks[-1]['text']='报名资格，2026年10月18日前提交材料。'
    selected=select_model_blocks(blocks)
    assert selected['characters']<=24000 and selected['omitted_count']>180 and any(b['page']==200 for b in selected['blocks'])
    assert len(blocks)==200 and len(blocks[0]['text'])==2495


def test_docx_decompression_bomb_rejected(tmp_path):
    path=tmp_path/'bomb.docx'
    with zipfile.ZipFile(path,'w',compression=zipfile.ZIP_DEFLATED) as z:z.writestr('[Content_Types].xml','ok');z.writestr('word/document.xml','a'*2000000)
    with pytest.raises(DocumentError) as e:validate_file(path,'docx')
    assert e.value.code=='decompression_limit'


@pytest.mark.parametrize('target',['http://127.0.0.1/x','https://evil.example/x','file:///C:/x'])
def test_redirect_outside_allowlist_rejected(tmp_path,target):
    with pytest.raises(DocumentError):download('https://www.zqu.edu.cn/a.pdf',['www.zqu.edu.cn'],tmp_path/'remote',transport=httpx.MockTransport(lambda _:httpx.Response(302,headers={'location':target})),resolver=lambda _:None)


def test_upload_async_parser_and_restart(store,tmp_path):
    service=DocumentService(store);app=FastAPI();app.include_router(router);app.include_router(review_router);app.add_middleware(UploadBodyLimit);app.state.document_service=service;app.state.personal_store=store
    try:
        with TestClient(app) as client:
            raw=client.post('/api/personal/collect-files',files={'files':('notice.txt','2026年10月18日前提交材料。'.encode(),'text/plain')},data={'title_hint':'上传验收'}).json()
            job=raw['document_job'];deadline=time.monotonic()+20
            while time.monotonic()<deadline:
                value=service.job(job['id'])
                if value['status'] not in ('running','waiting'):break
                time.sleep(.1)
            assert value['status']=='completed',value
            result=ReviewStore(store).get(value['result_snapshot_id']);assert result['snapshot']['blocks']
            b=result['snapshot']['blocks'][0];event=ReviewStore(store).confirm(result['snapshot']['id'],review_payload(result,b,'2026年10月18日前'))['events'][0]
    finally:service.close()
    reopened=PersonalStore(store.path);reopened.initialize();assert reopened.get_event(event['id'])==event


def test_failed_job_saved_original_retry_state(store,tmp_path):
    service=DocumentService(store)
    try:
        path=tmp_path/'bad.pdf';path.write_bytes(b'MZ bad data');raw=service.files.collect([(path.name,'application/pdf',path)],'','坏 PDF')
        job=service.start(raw['snapshot']['id'],[raw['snapshot']['attachments'][0]['id']]);end=time.monotonic()+10
        while time.monotonic()<end:
            value=service.job(job['id'])
            if value['status']=='failed':break
            time.sleep(.1)
        assert value['status']=='failed' and service.files.get(raw['snapshot']['attachments'][0]['id'])['parse_status']=='failed'
        assert service.files.path(service.files.get(raw['snapshot']['attachments'][0]['id'])).is_file()
    finally:service.close()


def test_same_name_files_in_one_collection_are_not_discarded(store,tmp_path):
    a=tmp_path/'a';b=tmp_path/'b';a.mkdir();b.mkdir()
    p=a/'notice.txt';q=b/'notice.txt';p.write_text('附件 A：2026年10月18日前提交。',encoding='utf-8');q.write_text('附件 B：2026年10月19日前提交。',encoding='utf-8')
    files=AttachmentStore(store);raw=files.collect([(p.name,'text/plain',p),(q.name,'text/plain',q)],'正文','同名两份文件')
    for att,path in zip(raw['snapshot']['attachments'],[p,q]):
        result=parse_document(path,'txt',tmp_path)
        with store.connection(write=True) as c:c.execute("UPDATE attachment SET parse_status='parsed',parse_manifest_json=?,extractor_version='documents-1' WHERE id=?",(encode(result),att['id']))
    snapshot=ReviewStore(store).get(files.compose(raw['snapshot']['id'],[a['id'] for a in raw['snapshot']['attachments']]))['snapshot']
    assert len(snapshot['attachments'])==2 and len([b for b in snapshot['blocks'] if b['attachment_id']])==2


def test_unknown_length_upload_stream_limit():
    received=[{'type':'http.request','body':b'x'*30_000_001,'more_body':True},{'type':'http.request','body':b'x'*30_000_001,'more_body':False}];sent=[]
    async def endpoint(*_):raise AssertionError('oversized request reached the multipart parser')
    async def receive():return received.pop(0)
    async def send(message):sent.append(message)
    asyncio.run(UploadBodyLimit(endpoint)({'type':'http','method':'POST','path':'/api/personal/collect-files','headers':[]},receive,send))
    assert sent[0]['status']==413


def test_equal_offsets_across_pages_remain_distinct(store,tmp_path):
    raw=local_snapshot(store,text_pdf(tmp_path/'repeat.pdf',['Submit application before 2026-10-18.']*2))
    blocks=raw['snapshot']['blocks'];payload=review_payload(raw,blocks[0],'2026-10-18').model_dump()
    payload['events'][0]['deadlines'][0]['evidence'].append({**payload['events'][0]['deadlines'][0]['evidence'][0],'block_id':blocks[1]['block_id']})
    event=ReviewStore(store).confirm(raw['snapshot']['id'],ConfirmationInput.model_validate(payload))['events'][0]
    assert len([e for e in event['evidence'] if e['field_path'].startswith('deadlines.')])==2


def test_file_failure_keeps_body_reviewable_with_warning(store,tmp_path):
    service=DocumentService(store)
    try:
        path=tmp_path/'bad.pdf';path.write_bytes(b'MZ bad data');raw=service.files.collect([(path.name,'application/pdf',path)],'2026年10月18日前提交材料。','正文可整理')
        job=service.start(raw['snapshot']['id'],[raw['snapshot']['attachments'][0]['id']]);end=time.monotonic()+10
        while time.monotonic()<end:
            value=service.job(job['id'])
            if value['status'] not in ('running','waiting'):break
            time.sleep(.1)
        assert value['status']=='partial' and value['result_snapshot_id']
        result=ReviewStore(store).get(value['result_snapshot_id']);assert result['snapshot']['blocks'][0]['text']=='2026年10月18日前提交材料。'
        assert result['snapshot']['document_warnings'][0]['code']=='files_unparsed'
        assert service.files.path(service.files.get(result['snapshot']['attachments'][0]['id'])).is_file()
    finally:service.close()


def test_interrupted_job_recovers_as_visible_retryable_error(store,tmp_path):
    files=AttachmentStore(store);path=tmp_path/'notice.txt';path.write_text('原文保留。',encoding='utf-8');raw=files.collect([(path.name,'text/plain',path)],'','中断验收');sid=raw['snapshot']['id'];aid=raw['snapshot']['attachments'][0]['id']
    with store.connection(write=True) as c:
        c.execute("UPDATE attachment SET parse_status='parsing' WHERE id=?",(aid,))
        c.execute("INSERT INTO attachment_job (id,raw_item_id,source_snapshot_id,attachment_ids_json,operation,status,phase,created_at) VALUES ('interrupted',?,?,?,'parse','running','解析中','2026-10-07')",(raw['snapshot']['raw_item_id'],sid,encode([aid])))
    service=DocumentService(store)
    try:
        result=ReviewStore(store).get(sid)
        assert result['snapshot']['attachment_jobs'][0]['error'][0]['code']=='interrupted'
        assert result['snapshot']['attachments'][0]['parse_status']=='failed'
        assert files.path(files.get(aid)).is_file()
    finally:service.close()


def test_official_txt_images_and_filename_without_suffix_are_discovered():
    from server.official_sources import load_sources
    from server.official_parsers import parse_article
    source=next(s for s in load_sources() if s.id=='zqu-school')
    html='<div class="v_news_content">具体安排见附件。'+''.join(f'<a href="/files/notice.{ext}">附件{ext}</a>' for ext in ('pdf','docx','txt','png','jpg','jpeg'))+'</div>'
    article=parse_article(html.encode(),'https://www.zqu.edu.cn/info/1361/1.htm',source,{'title':'通知','published_at':None})
    assert {a['extension'] for a in article['attachments']}=={'pdf','docx','txt','png','jpg','jpeg'}


def test_downloaded_format_can_be_derived_for_suffixless_official_name(store,tmp_path):
    path=text_pdf(tmp_path/'document.pdf',['Submit application before 2026-10-18.'])
    files=AttachmentStore(store);raw=files.collect([('附件通知',None,path)],'正文','无后缀链接');att=raw['snapshot']['attachments'][0]
    downloaded=files.downloaded(att['id'],path,'application/pdf','pdf')
    assert downloaded['extension']=='pdf' and files.path(downloaded).suffix=='.pdf'
    assert parse_document(files.path(downloaded),'pdf',tmp_path)['blocks'][0]['page']==1


def test_redownload_hash_cache_hit_does_not_bypass_mime_policy(store,tmp_path,monkeypatch):
    from server.official_store import OfficialStore
    from server.official_sources import load_sources
    OfficialStore(store).configure(load_sources())
    path=text_pdf(tmp_path/'notice.pdf',['Submit application before 2026-10-18.']);raw=local_snapshot(store,path,url='https://www.zqu.edu.cn/notice.pdf')
    with store.connection(write=True) as c:c.execute("UPDATE raw_item SET source_id='zqu-school' WHERE id=?",(raw['snapshot']['raw_item_id'],))
    def fake(url,hosts,destination,**kwargs):Path(destination).write_bytes(path.read_bytes());return 'text/plain'
    monkeypatch.setattr('server.documents.service.download',fake)
    service=DocumentService(store)
    try:
        att=raw['snapshot']['attachments'][0];job=service.start(raw['snapshot']['id'],[att['id']],redownload=True);end=time.monotonic()+10
        while time.monotonic()<end:
            value=service.job(job['id'])
            if value['status'] not in ('running','waiting'):break
            time.sleep(.1)
        assert value['error'][0]['code']=='mime_mismatch'
        assert service.files.get(att['id'])['parse_status']=='parsed'
        assert service.files.path(service.files.get(att['id'])).read_bytes()==path.read_bytes()
    finally:service.close()
